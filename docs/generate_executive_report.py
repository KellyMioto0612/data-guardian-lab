"""Generate the Word version of the Data Guardian executive inventory."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).with_name("Data_Guardian_Relatorio_Executivo.docx")


def heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(14)
    paragraph.paragraph_format.space_after = Pt(6)


def bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def main() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)

    style = document.styles["Normal"]
    style.font.name = "Aptos"
    style.font.size = Pt(10.5)

    title = document.add_heading("Data Guardian", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Relatório executivo — inventário de entregas")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.color.rgb = RGBColor(89, 89, 89)

    heading(document, "Resumo executivo")
    document.add_paragraph(
        "O Data Guardian possui uma base funcional para observabilidade e governança "
        "técnica de dados: descobre ativos SQL, identifica riscos, calcula dívida "
        "técnica, mede impacto de dependências e prioriza ações. A operação Synapse "
        "foi implementada em modo somente leitura e aguarda a homologação controlada "
        "no workspace autorizado."
    )

    heading(document, "Status do roadmap")
    status = [
        ("Fundação", "██████████", "Concluída"),
        ("Descoberta SQL", "██████████", "Concluída"),
        ("Inteligência", "██████████", "Concluída no MVP"),
        ("Operação Synapse", "██████░░░░", "Implementada e aguardando homologação"),
        ("Produto e IA", "░░░░░░░░░░", "Futuro"),
    ]
    table = document.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    for cell, text in zip(table.rows[0].cells, ("Frente", "Progresso", "Situação"), strict=True):
        cell.text = text
    for name, progress, situation in status:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = name, progress, situation

    heading(document, "Fundação e produto")
    bullets(document, [
        "Modelos e contratos estáveis para pipelines, objetos e observações.",
        "Arquitetura desacoplada por provedores, com DemoProvider e SynapseProvider.",
        "Dataset demonstrativo e GuardianScout independente da origem dos dados.",
        "Dashboard Streamlit com KPIs, gráficos e filtros.",
    ])

    heading(document, "Descoberta e análise SQL")
    bullets(document, [
        "Scanner SQL com controle de diretório permitido e tratamento para T-SQL/Synapse.",
        "Inventário de procedures e views, com leitura de dependências e grafo de lineage.",
        "Proteção para não expor SQL bruto em mensagens de erro.",
        "Identificação de ativos persistentes, separando artefatos temporários e comandos auxiliares.",
    ])

    heading(document, "Inteligência técnica")
    bullets(document, [
        "Evidências para SELECT *, CROSS JOIN, JOIN sem condição, EXEC dinâmico, alta complexidade, subqueries, tabelas temporárias, falta de documentação e alterações sem limite.",
        "TDI (índice de dívida técnica) de 0 a 100, com classificação de severidade e criticidade.",
        "DNA estrutural e similaridade entre procedures e views.",
        "Impacto upstream/downstream, detecção de ciclos e ranking de prioridade técnica.",
    ])

    heading(document, "Segurança")
    bullets(document, [
        "Não há credenciais hardcoded: a configuração usa variáveis de ambiente e .env.example sem segredos.",
        "Autenticação Azure por Azure CLI, Managed Identity ou Workload Identity.",
        "Gitignore reforçado e quarentena privada para SQLs brutos, inventários e backups fora do projeto.",
        "Erros de autenticação, rede e parser são tratados sem vazar token, conteúdo SQL ou dados sensíveis.",
    ])

    heading(document, "Operação Synapse")
    bullets(document, [
        "Conector implementado exclusivamente para leitura de execuções de pipeline.",
        "Janela de consulta limitada de 1 a 168 horas, validação de workspace e paginação limitada.",
        "Estados de execução normalizados para o modelo interno e integração ao dashboard.",
        "A primeira consulta real ainda não foi executada após o logout do Azure CLI; nenhuma sessão do Guardian permanece ativa.",
    ])

    heading(document, "Qualidade e validação")
    bullets(document, [
        "108 testes automatizados aprovados; 1 teste opcional ignorado.",
        "Lint sem erros e verificação de dependências sem conflitos.",
        "Testes do conector Synapse usam simulação de transporte, sem consulta ao Azure.",
    ])

    heading(document, "Próximos passos")
    bullets(document, [
        "Homologar a operação Synapse com uma consulta real, autorizada e somente leitura.",
        "Adicionar métricas Azure Monitor/Synapse e telemetria segura do Guardian.",
        "Definir persistência aprovada, retenção, regras de SLA e alertas.",
        "Formalizar RBAC de produção e migrar execuções agendadas para Managed Identity.",
    ])

    document.add_paragraph("Documento gerado a partir do estado atual do repositório.").runs[0].italic = True
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
